#!/usr/bin/env python3
"""インスタの予約投稿用に、画像とキャプションを1か所にまとめる。

使い方:
    python3 prep-instagram.py          # 次の5本ぶん
    python3 prep-instagram.py 8        # 次の8本ぶん

~/Downloads/インスタ予約_YYYY-MM-DD/ に、投稿日ごとのフォルダを作る。
各フォルダの中身:
    1_アイキャッチ.png   … ブログのアイキャッチ
    2_<名前>.png        … 表などの追加画像（あれば）
    キャプション.txt      … xlsxのキャプションをそのまま
    _情報.txt            … 投稿日・記事タイトル・ブログURL
"""
import openpyxl, glob, re, os, shutil, sys, unicodedata
from datetime import date

BLOG = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BLOG, 'インスタ投稿スケジュール_ドタバタ母さんブログ.xlsx')
OUT_ROOT = os.path.expanduser('~/Downloads')
EXTRA_DIR = os.path.expanduser('~/Downloads')   # 表の画像などを置いてある場所

N = int(sys.argv[1]) if len(sys.argv) > 1 else 5


def norm(s):
    """全角/半角や記号ゆれを吸収してタイトルを突き合わせる"""
    s = unicodedata.normalize('NFKC', s or '')
    return re.sub(r'\s+', '', s)


# 記事タイトル → (スラッグ, アイキャッチの実ファイル)
articles = {}
for path in glob.glob(os.path.join(BLOG, 'content/posts/*.md')):
    src = open(path, encoding='utf-8').read()
    head = src.split('---', 2)[1] if src.startswith('---') else src[:1200]
    t = re.search(r'^title:\s*"(.*)"', head, re.M)
    c = re.search(r'^\s*image:\s*"(.*)"', head, re.M)
    if not t:
        continue
    slug = os.path.basename(path)[:-3]
    img = None
    if c:
        p = os.path.join(BLOG, 'static', c.group(1).lstrip('/'))
        if os.path.exists(p):
            img = p
    articles[norm(t.group(1))] = (slug, img, t.group(1))

ws = openpyxl.load_workbook(XLSX).active
pending = []
for r in range(2, ws.max_row + 1):
    if '予定' not in (ws.cell(r, 5).value or ''):
        continue
    pending.append({
        'date': ws.cell(r, 4).value,
        'title': ws.cell(r, 2).value or '',
        'caption': ws.cell(r, 6).value or '',
    })
pending = pending[:N]

if not pending:
    print('予定の投稿がありません')
    sys.exit()

out = os.path.join(OUT_ROOT, 'インスタ予約_' + date.today().isoformat())
os.makedirs(out, exist_ok=True)

missing = []
for i, p in enumerate(pending, 1):
    slug, img, real_title = articles.get(norm(p['title']), (None, None, None))
    d = p['date'].replace('/', '')[4:]          # MMDD
    name = f"{i:02d}_{d}_{(slug or 'unknown')}"
    folder = os.path.join(out, name)
    os.makedirs(folder, exist_ok=True)

    if img:
        shutil.copy(img, os.path.join(folder, '1_アイキャッチ' + os.path.splitext(img)[1]))
    else:
        missing.append(f"{p['date']}  アイキャッチなし  {p['title'][:34]}")

    # 追加画像（Downloads の「インスタ_*.png」から、スラッグが名前に入っているもの）
    n = 2
    if slug:
        for extra in sorted(glob.glob(os.path.join(EXTRA_DIR, 'インスタ_*.png'))):
            if slug in os.path.basename(extra):
                shutil.copy(extra, os.path.join(folder, f'{n}_' + os.path.basename(extra)))
                n += 1

    with open(os.path.join(folder, 'キャプション.txt'), 'w', encoding='utf-8') as f:
        f.write(p['caption'])

    with open(os.path.join(folder, '_情報.txt'), 'w', encoding='utf-8') as f:
        f.write(f"投稿日: {p['date']} 10:00\n")
        f.write(f"記事: {real_title or p['title']}\n")
        if slug:
            f.write(f"URL : https://3nin-dotabata.com/posts/{slug}/\n")

print(f'{len(pending)}本ぶんを用意しました')
print(out)
if missing:
    print('\n--- 画像が見つからなかったもの（手で入れてください）---')
    for m in missing:
        print(' ', m)
